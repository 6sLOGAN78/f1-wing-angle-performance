#!/usr/bin/env python3
"""Comprehensive cross-artifact verification of the F1 Wing Angle Performance project."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

# Support running directly or as a module
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
from f1wing.presentation import all_slide_text_within_safe_bounds

# Import MATLAB static contract checker
sys.path.insert(0, str(Path(__file__).resolve().parent))
from check_matlab_static import check_matlab_tree

REQUIRED_REPORT_HEADINGS = [
    "Abstract",
    "1. Introduction and Research Question",
    "2. Literature Review",
    "3. Aerodynamic Theory",
    "4. Vehicle and Tyre Dynamics",
    "5. Model Architecture and Implementation",
    "6. Verification and Reproducibility",
    "7. Results",
    "8. Uncertainty and Sensitivity",
    "9. Engineering Interpretation",
    "10. Assumptions and Limitations",
    "11. Conclusions and Future Work",
    "References",
    "Appendix A — Parameter Register",
    "Appendix B — MATLAB and Simulink Workflow",
]


@dataclass
class VerificationReport:
    run_id: str
    run_id_consistent: bool
    headline_values_consistent: bool
    missing_artifacts: list[str]
    hash_mismatches: list[str]
    report_valid: bool
    presentation_valid: bool
    matlab_valid: bool
    all_passed: bool
    details: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def _compute_sha256(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()


def _get_docx_text(path: Path) -> str:
    doc = Document(str(path))
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def _get_pptx_text(path: Path) -> str:
    prs = Presentation(str(path))
    chunks = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        chunks.append(cell.text)
    return "\n".join(chunks)


def verify_project(project_root: Path) -> VerificationReport:
    """Run all cross-artifact consistency, integrity, and reproducibility assertions."""
    root = Path(project_root).resolve()
    manifest_path = root / "results" / "manifest.json"
    if not manifest_path.is_file():
        raise FileNotFoundError(f"Manifest not found at {manifest_path}")

    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    run_id = manifest.get("run_id", "")
    headline_metrics = manifest.get("headline_metrics", {})

    missing_artifacts: list[str] = []
    hash_mismatches: list[str] = []
    details: dict[str, Any] = {
        "checked_inputs": {},
        "checked_artifacts": {},
        "report_checks": {},
        "presentation_checks": {},
        "matlab_checks": {},
    }

    # 1. Verify input files against manifest hashes
    for rel_path, expected_hash in manifest.get("input_sha256", {}).items():
        file_path = root / rel_path
        if not file_path.is_file():
            missing_artifacts.append(rel_path)
            details["checked_inputs"][rel_path] = "MISSING"
        else:
            actual_hash = _compute_sha256(file_path)
            if actual_hash != expected_hash:
                hash_mismatches.append(rel_path)
                details["checked_inputs"][rel_path] = f"MISMATCH (got {actual_hash[:8]}, expected {expected_hash[:8]})"
            else:
                details["checked_inputs"][rel_path] = "MATCH"

    # 2. Verify artifact files against manifest hashes
    for rel_path, expected_hash in manifest.get("artifact_sha256", {}).items():
        file_path = root / "results" / rel_path
        if not file_path.is_file():
            missing_artifacts.append(str(Path("results") / rel_path))
            details["checked_artifacts"][rel_path] = "MISSING"
        else:
            actual_hash = _compute_sha256(file_path)
            if actual_hash != expected_hash:
                hash_mismatches.append(str(Path("results") / rel_path))
                details["checked_artifacts"][rel_path] = f"MISMATCH (got {actual_hash[:8]}, expected {expected_hash[:8]})"
            else:
                details["checked_artifacts"][rel_path] = "MATCH"

    # 3. Verify Report (DOCX & PDF)
    report_docx = root / "results" / "report" / "F1_Wing_Angle_Performance_Report.docx"
    report_pdf = root / "results" / "report" / "F1_Wing_Angle_Performance_Report.pdf"

    report_valid = True
    report_run_id_found = False
    report_headline_found = True

    if not report_docx.is_file():
        missing_artifacts.append(str(report_docx.relative_to(root)))
        report_valid = False
    if not report_pdf.is_file():
        missing_artifacts.append(str(report_pdf.relative_to(root)))
        report_valid = False

    if report_docx.is_file():
        report_text = _get_docx_text(report_docx)
        report_run_id_found = run_id in report_text

        # Check required headings
        missing_headings = [h for h in REQUIRED_REPORT_HEADINGS if h not in report_text]
        details["report_checks"]["missing_headings"] = missing_headings
        if missing_headings:
            report_valid = False

        # Check ethical disclosure
        details["report_checks"]["educational_disclosure"] = "representative educational model" in report_text.lower()
        if not details["report_checks"]["educational_disclosure"]:
            report_valid = False

        # Check headline metrics
        missing_report_metrics = []
        for grp_name, grp in headline_metrics.items():
            for c_name, val in grp.items():
                formatted = f"{val:.3f}"
                if formatted not in report_text:
                    missing_report_metrics.append(f"{grp_name}.{c_name}={formatted}")
        details["report_checks"]["missing_headline_metrics"] = missing_report_metrics
        if missing_report_metrics:
            report_headline_found = False
            report_valid = False

    # 4. Verify Presentation (PPTX & PDF)
    pres_pptx = root / "results" / "presentation" / "F1_Wing_Angle_Performance_Presentation.pptx"
    pres_pdf = root / "results" / "presentation" / "F1_Wing_Angle_Performance_Presentation.pdf"

    presentation_valid = True
    pres_run_id_found = False
    pres_headline_found = True

    if not pres_pptx.is_file():
        missing_artifacts.append(str(pres_pptx.relative_to(root)))
        presentation_valid = False
    if not pres_pdf.is_file():
        missing_artifacts.append(str(pres_pdf.relative_to(root)))
        presentation_valid = False

    if pres_pptx.is_file():
        deck = Presentation(str(pres_pptx))
        slide_count = len(deck.slides)
        details["presentation_checks"]["slide_count"] = slide_count
        if not (16 <= slide_count <= 20):
            presentation_valid = False

        safe_bounds = all_slide_text_within_safe_bounds(deck)
        details["presentation_checks"]["safe_bounds"] = safe_bounds
        if not safe_bounds:
            presentation_valid = False

        pres_text = _get_pptx_text(pres_pptx)
        pres_run_id_found = run_id in pres_text

        # Check headline metrics
        missing_pres_metrics = []
        for grp_name, grp in headline_metrics.items():
            for c_name, val in grp.items():
                formatted = f"{val:.3f}"
                if formatted not in pres_text:
                    missing_pres_metrics.append(f"{grp_name}.{c_name}={formatted}")
        details["presentation_checks"]["missing_headline_metrics"] = missing_pres_metrics
        if missing_pres_metrics:
            pres_headline_found = False
            presentation_valid = False

    # 5. Verify MATLAB static compliance & README disclosure
    matlab_report = check_matlab_tree(root / "matlab")
    matlab_valid = matlab_report.ok
    details["matlab_checks"]["missing_files"] = matlab_report.missing_files
    details["matlab_checks"]["signature_errors"] = matlab_report.signature_errors
    details["matlab_checks"]["forbidden_paths"] = matlab_report.forbidden_absolute_paths

    readme_path = root / "README.md"
    readme_text = readme_path.read_text(encoding="utf-8") if readme_path.is_file() else ""
    matlab_disclosed = "matlab" in readme_text.lower() and "r2022b" in readme_text.lower()
    details["matlab_checks"]["runtime_limitation_disclosed"] = matlab_disclosed
    if not matlab_disclosed:
        matlab_valid = False

    # 6. Synthesize consistency & overall status
    run_id_consistent = bool(run_id and report_run_id_found and pres_run_id_found)
    headline_values_consistent = bool(report_headline_found and pres_headline_found)

    all_passed = (
        len(missing_artifacts) == 0
        and len(hash_mismatches) == 0
        and run_id_consistent
        and headline_values_consistent
        and report_valid
        and presentation_valid
        and matlab_valid
    )

    report_obj = VerificationReport(
        run_id=run_id,
        run_id_consistent=run_id_consistent,
        headline_values_consistent=headline_values_consistent,
        missing_artifacts=missing_artifacts,
        hash_mismatches=hash_mismatches,
        report_valid=report_valid,
        presentation_valid=presentation_valid,
        matlab_valid=matlab_valid,
        all_passed=all_passed,
        details=details,
    )

    # Write out machine-readable verification report
    verification_dir = root / "results" / "verification"
    verification_dir.mkdir(parents=True, exist_ok=True)
    report_json = verification_dir / "verification_report.json"
    report_json.write_text(json.dumps(report_obj.to_dict(), indent=2), encoding="utf-8")

    return report_obj


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=Path("."))
    args = parser.parse_args()

    report = verify_project(args.project_root)
    print("=" * 70)
    print("F1 WING ANGLE PERFORMANCE: END-TO-END VERIFICATION REPORT")
    print("=" * 70)
    print(f"Run ID:                      {report.run_id}")
    print(f"Run ID Consistent:           {'PASS' if report.run_id_consistent else 'FAIL'}")
    print(f"Headline Values Consistent:  {'PASS' if report.headline_values_consistent else 'FAIL'}")
    print(f"Report Valid (DOCX & PDF):   {'PASS' if report.report_valid else 'FAIL'}")
    print(f"Presentation Valid (PPTX/PDF):{'PASS' if report.presentation_valid else 'FAIL'}")
    print(f"MATLAB Code & Static Checks: {'PASS' if report.matlab_valid else 'FAIL'}")
    print(f"Missing Artifacts:           {len(report.missing_artifacts)}")
    print(f"Hash Mismatches:             {len(report.hash_mismatches)}")
    print("-" * 70)
    print(f"OVERALL STATUS:              {'ALL CHECKS PASSED' if report.all_passed else 'VERIFICATION FAILED'}")
    print("=" * 70)

    if not report.all_passed:
        if report.missing_artifacts:
            print("Missing Artifacts:")
            for item in report.missing_artifacts:
                print(f"  - {item}")
        if report.hash_mismatches:
            print("Hash Mismatches:")
            for item in report.hash_mismatches:
                print(f"  - {item}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
