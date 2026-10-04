"""Generate publication-style project documents from the analysis manifest."""

from __future__ import annotations

import json
import html
import io
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image as PdfImage,
    LongTable,
    PageBreak,
    Paragraph as PdfParagraph,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)


NAVY = "13233A"
RED = "D71920"
LIGHT_BLUE = "EAF0F6"
MID_GREY = "64748B"
WHITE = "FFFFFF"


@dataclass
class ReportContext:
    project_root: Path
    results_root: Path
    manifest: dict
    document: Document
    table_number: int = 0


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, code, separate, end))


def _configure_document(document: Document, run_id: str) -> None:
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color in (
        ("Title", 30, NAVY),
        ("Subtitle", 15, MID_GREY),
        ("Heading 1", 20, NAVY),
        ("Heading 2", 14, RED),
        ("Heading 3", 11, NAVY),
    ):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if name == "Heading 1":
            style.paragraph_format.page_break_before = True
            style.paragraph_format.space_after = Pt(12)

    caption = document.styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MID_GREY)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if "Equation" not in [style.name for style in document.styles]:
        equation = document.styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
        equation.font.name = "Cambria Math"
        equation.font.size = Pt(11)
        equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        equation.paragraph_format.space_before = Pt(6)
        equation.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "ME3106 · F1 WING-ANGLE PERFORMANCE STUDY"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MID_GREY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"Run {run_id}  ·  PAGE ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)
    _add_field(footer, "PAGE")

    settings = document.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _add_cover(document: Document, manifest: dict) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(42)
    run = paragraph.add_run("ADVANCED MOTORSPORT SIMULATION")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(RED)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Simulation and Analysis of the\nEffect of Rear-Wing Angle\non F1 Car Performance")

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run("Aerodynamics · Vehicle Dynamics · Lap-Time Optimization · Active Aero")

    document.add_paragraph("\n")
    line = document.add_table(rows=1, cols=1)
    line.autofit = False
    line.columns[0].width = Inches(6.5)
    _set_cell_shading(line.cell(0, 0), RED)
    line.cell(0, 0).text = ""
    line.rows[0].height = Cm(0.12)

    facts = document.add_table(rows=0, cols=2)
    facts.alignment = WD_TABLE_ALIGNMENT.LEFT
    for label, value in (
        ("Project", manifest["model_name"]),
        ("Model version", manifest["model_version"]),
        ("Run ID", manifest["run_id"]),
        ("Generated", manifest["generated_at_utc"]),
        ("Classification", "Representative educational model"),
        ("Toolchain", "Python · MATLAB R2022b+ source · Simulink builder"),
    ):
        cells = facts.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
        cells[0].paragraphs[0].runs[0].bold = True
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(28)
    note_run = note.add_run(
        "Academic engineering report. Numerical values are traceable to the manifest and must "
        "be interpreted within the stated assumptions and limitations."
    )
    note_run.italic = True
    note_run.font.color.rgb = RGBColor.from_string(MID_GREY)
    document.add_page_break()


def _add_front_matter(document: Document, manifest: dict) -> None:
    document.add_heading("Document Control", level=1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    rows = (
        ("Title", "Simulation and Analysis of the Effect of Rear-Wing Angle on F1 Car Performance"),
        ("Purpose", "Reproducible advanced educational engineering study"),
        ("Run ID", manifest["run_id"]),
        ("Random seed", manifest["seed"]),
        ("Configuration", "config/project.json and config/uncertainty.json"),
        ("Evidence bundle", "results/manifest.json and SHA-256 artifact register"),
        ("MATLAB status", "R2022b+-compatible source; not runtime-tested in authoring environment"),
    )
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)
        cells[0].paragraphs[0].runs[0].bold = True

    document.add_heading("Contents", level=1)
    contents = (
        "Abstract",
        "1. Introduction and Research Question",
        "2. Literature Review",
        "3. Aerodynamic Theory",
        "4. Vehicle and Tyre Dynamics",
        "5. Model Architecture and Implementation",
        "6. Verification and Reproducibility",
        "7. Results",
        "8. Uncertainty and Sensitivity",
        "9. Engineering Interpretation",
        "10. Assumptions and Limitations",
        "11. Conclusions and Future Work",
        "References",
        "Appendices A–D",
    )
    for item in contents:
        paragraph = document.add_paragraph(item)
        paragraph.paragraph_format.left_indent = Cm(0.4)
        paragraph.paragraph_format.space_after = Pt(3)


def _replace_placeholders(text: str, manifest: dict) -> str:
    values = manifest["headline_metrics"]
    replacements = {
        "RUN_ID": manifest["run_id"],
        "SEED": str(manifest["seed"]),
        "OPT_LOW_ANGLE": f'{values["fixed_optimum_angle_deg"]["low_downforce"]:.3f}',
        "OPT_BAL_ANGLE": f'{values["fixed_optimum_angle_deg"]["balanced"]:.3f}',
        "OPT_HIGH_ANGLE": f'{values["fixed_optimum_angle_deg"]["high_downforce"]:.3f}',
        "OPT_LOW_TIME": f'{values["fixed_minimum_lap_time_s"]["low_downforce"]:.3f}',
        "OPT_BAL_TIME": f'{values["fixed_minimum_lap_time_s"]["balanced"]:.3f}',
        "OPT_HIGH_TIME": f'{values["fixed_minimum_lap_time_s"]["high_downforce"]:.3f}',
        "ACTIVE_LOW_TIME": f'{values["active_limited_lap_time_s"]["low_downforce"]:.3f}',
        "ACTIVE_BAL_TIME": f'{values["active_limited_lap_time_s"]["balanced"]:.3f}',
        "ACTIVE_HIGH_TIME": f'{values["active_limited_lap_time_s"]["high_downforce"]:.3f}',
    }
    for key, value in replacements.items():
        text = text.replace("{" + key + "}", value)
    return text


def _format_value(value) -> str:
    if isinstance(value, (list, tuple)):
        return "–".join(_format_value(item) for item in value)
    if pd.isna(value):
        return "—"
    if isinstance(value, float):
        magnitude = abs(value)
        if magnitude >= 1000:
            return f"{value:,.1f}"
        if magnitude >= 100:
            return f"{value:.2f}"
        return f"{value:.3f}"
    return str(value).replace("_", " ")


def _add_table(ctx: ReportContext, title: str, headers: Iterable[str], rows: Iterable[Iterable]) -> None:
    ctx.table_number += 1
    caption = ctx.document.add_paragraph(style="Caption")
    caption.add_run(f"Table {ctx.table_number}. {title}")
    headers = list(headers)
    table = ctx.document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        _set_cell_shading(cell, NAVY)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
            run.font.size = Pt(8)
    _set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = _format_value(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)
        if row_index % 2:
            for cell in cells:
                _set_cell_shading(cell, LIGHT_BLUE)


def _dataframe(ctx: ReportContext, name: str) -> pd.DataFrame:
    return pd.read_csv(ctx.results_root / "data" / f"{name}.csv")


def _table_content(ctx: ReportContext, key: str) -> None:
    config = json.loads((ctx.project_root / "config" / "project.json").read_text())
    if key == "project_scope":
        _add_table(
            ctx,
            "Study scope and outputs.",
            ("Item", "Definition"),
            (
                ("Independent variable", "Rear-wing angle, 0–30 deg"),
                ("Circuit cases", "Low-downforce, balanced, high-downforce synthetic closed tracks"),
                ("Strategies", "Fixed, open-wing, ideal-active, actuator-limited active"),
                ("Primary response", "Lap time with straight-line, braking, cornering, energy diagnostics"),
                ("Uncertainty", "1,000 bounded samples per circuit; seed 3106"),
                ("Deliverables", "CSV, XLSX, MAT, PNG, SVG, DOCX, PDF, PPTX, MATLAB and Simulink builder"),
            ),
        )
    elif key == "nomenclature":
        _add_table(
            ctx,
            "Principal nomenclature.",
            ("Symbol", "Meaning", "Unit"),
            (
                ("A", "Full-car aerodynamic reference area", "m²"),
                ("AR", "Rear-wing aspect ratio", "—"),
                ("C_L", "Positive downforce coefficient magnitude", "—"),
                ("C_D", "Drag coefficient", "—"),
                ("e", "Span-efficiency factor", "—"),
                ("F_z", "Normal load", "N"),
                ("mu", "Tyre friction coefficient", "—"),
                ("q", "Dynamic pressure", "Pa"),
                ("rho", "Air density", "kg/m³"),
                ("v", "Vehicle speed", "m/s"),
            ),
        )
    elif key == "tracks":
        rows = []
        for name in ("low_downforce", "balanced", "high_downforce"):
            track = json.loads((ctx.project_root / "tracks" / f"{name}.json").read_text())
            rows.append(
                (
                    name,
                    sum(float(segment["length_m"]) for segment in track["segments"]),
                    len(track["segments"]),
                    max(abs(float(segment["curvature_start_1pm"])) for segment in track["segments"]),
                    sum(bool(segment["aero_eligible"]) for segment in track["segments"]),
                )
            )
        _add_table(ctx, "Synthetic circuit definitions.", ("Track", "Length (m)", "Segments", "Max |curvature| (1/m)", "Aero-eligible segments"), rows)
    elif key == "optima":
        frame = _dataframe(ctx, "circuit_optima")
        cols = ["track", "optimum_angle_deg", "minimum_lap_time_s", "neighbour_sensitivity_s_per_deg", "optimization_method"]
        _add_table(ctx, "Optimized fixed-wing results.", ("Track", "Angle (deg)", "Lap time (s)", "Neighbour sensitivity (s/deg)", "Method"), frame[cols].itertuples(index=False, name=None))
    elif key == "strategies":
        frame = _dataframe(ctx, "strategy_comparison")
        cols = ["track", "mode", "open_angle_deg", "closed_angle_deg", "lap_time_s", "maximum_speed_kph", "tractive_energy_mj"]
        _add_table(ctx, "Aerodynamic strategy comparison.", ("Track", "Mode", "Open (deg)", "Closed (deg)", "Lap (s)", "Vmax (km/h)", "Wheel work (MJ)"), frame[cols].itertuples(index=False, name=None))
    elif key == "uncertainty":
        frame = _dataframe(ctx, "monte_carlo_summary")
        frame = frame[frame["metric"].isin(["optimum_angle_deg", "minimum_lap_time_s"])]
        cols = ["track", "percentile", "metric", "value"]
        _add_table(ctx, "Monte Carlo percentile summary.", ("Track", "Percentile", "Metric", "Value"), frame[cols].itertuples(index=False, name=None))
    elif key == "sensitivity":
        frame = _dataframe(ctx, "sensitivity")
        cols = ["parameter", "perturbation_percent", "lap_time_s", "change_from_nominal_s"]
        _add_table(ctx, "Balanced-circuit local sensitivity.", ("Parameter", "Change (%)", "Lap (s)", "Delta (s)"), frame[cols].itertuples(index=False, name=None))
    elif key == "verification":
        artifacts = ctx.manifest["artifact_sha256"]
        rows = (
            ("Input provenance", len(ctx.manifest["input_sha256"]), "SHA-256 registered"),
            ("Output artifacts", len(artifacts), "SHA-256 registered"),
            ("Numerical tables", len(ctx.manifest["tables"]), "CSV + XLSX mapping"),
            ("Figures", len(ctx.manifest["figures"]), "PNG + SVG + source CSV"),
            ("Randomness", 1, f'Deterministic seed {ctx.manifest["seed"]}'),
            ("MATLAB parity", 12, "Static contract + reference fixtures"),
        )
        _add_table(ctx, "Verification evidence summary.", ("Check", "Count", "Evidence"), rows)
    elif key == "limitations":
        limitation_text = (ctx.project_root / "report" / "assumptions_and_limitations.md").read_text()
        bullets = [line[2:] for line in limitation_text.splitlines() if line.startswith("- ")]
        _add_table(ctx, "Assumptions and limitations register.", ("ID", "Declared limitation"), ((f"L{index:02d}", item) for index, item in enumerate(bullets, start=1)))
    elif key == "parameters":
        rows = []
        for name, details in config["parameter_register"].items():
            rows.append((name, details["unit"], details["classification"], details["supported_range"], details["description"]))
        _add_table(ctx, "Model parameter register.", ("Parameter", "Unit", "Class", "Supported range", "Description"), rows)
    elif key == "software":
        _add_table(
            ctx,
            "Software workflow and runtime status.",
            ("Layer", "Purpose", "Status"),
            (
                ("Python 3.10+", "Reference analysis, optimization, uncertainty, exports", "Runtime-tested"),
                ("MATLAB R2022b+", "Independent numerical mirror and parity assertions", "Static-checked; runtime pending"),
                ("Simulink", "Generated actuator/aero/longitudinal demonstration", "Builder supplied; license required"),
                ("LibreOffice", "DOCX/PPTX to PDF conversion", "Runtime-tested"),
                ("pytest", "Physics, contract, export, and document checks", "Automated"),
            ),
        )
    elif key == "artifacts":
        rows = []
        for path, digest in sorted(ctx.manifest["artifact_sha256"].items()):
            rows.append((path, digest[:16] + "…"))
        _add_table(ctx, "Manifest artifact index (abbreviated hashes).", ("Artifact", "SHA-256 prefix"), rows)
    else:
        raise KeyError(f"Unknown report table token: {key}")


def _add_figure(ctx: ReportContext, figure_id: str) -> None:
    figures = ctx.manifest["figures"]
    matches = [(index, value) for index, value in enumerate(figures, start=1) if value["id"] == figure_id]
    if not matches:
        raise KeyError(f"Unknown figure token: {figure_id}")
    number, figure = matches[0]
    path = ctx.results_root / figure["path"]
    if not path.is_file():
        raise FileNotFoundError(path)
    ctx.document.add_page_break()
    heading = ctx.document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = heading.add_run(f"RESULT PLATE {number:02d}")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(RED)
    picture_paragraph = ctx.document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = picture_paragraph.add_run().add_picture(str(path), width=Inches(6.55))
    properties = picture._inline.docPr
    properties.set("descr", figure["caption"])
    caption = ctx.document.add_paragraph(style="Caption")
    caption.add_run(f"Figure {number}. {figure['caption']}")
    source = ctx.document.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source_run = source.add_run(
        f"Source data: {figure['source_data']} · Units: {figure['units']} · Run: {ctx.manifest['run_id']}"
    )
    source_run.font.size = Pt(7.5)
    source_run.font.color.rgb = RGBColor.from_string(MID_GREY)


def _add_equation(document: Document, number: int, expression: str) -> None:
    paragraph = document.add_paragraph(style="Equation")
    paragraph.add_run(expression.replace("_", "₋"))
    label = document.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = label.add_run(f"Equation ({number})")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GREY)


def _add_references(ctx: ReportContext) -> None:
    entries = json.loads((ctx.project_root / "references" / "references.json").read_text())
    for entry in entries:
        parts = [
            f"[{entry['id']}] {entry['authors']} ({entry['year']}).",
            f"{entry['title']}.",
            f"{entry['publisher']}.",
        ]
        if entry.get("doi"):
            parts.append(f"doi:{entry['doi']}.")
        if entry.get("url"):
            parts.append(entry["url"])
        if entry.get("accessed"):
            parts.append(f"Accessed {entry['accessed']}.")
        paragraph = ctx.document.add_paragraph(" ".join(parts))
        paragraph.paragraph_format.first_line_indent = Cm(-0.6)
        paragraph.paragraph_format.left_indent = Cm(0.6)
        paragraph.paragraph_format.space_after = Pt(5)


TOKEN = re.compile(r"^\{\{([A-Z_]+)(?::([^}|]+))?(?:\|(.+))?\}\}$")


def _render_markdown(ctx: ReportContext, text: str) -> None:
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            ctx.document.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith("## "):
            ctx.document.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith("### "):
            ctx.document.add_heading(line[4:].strip(), level=3)
            continue
        if line.startswith("- "):
            ctx.document.add_paragraph(line[2:].strip(), style="List Bullet")
            continue
        token = TOKEN.match(line)
        if token:
            kind, argument, payload = token.groups()
            if kind == "FIGURE" and argument:
                _add_figure(ctx, argument)
            elif kind == "TABLE" and argument:
                _table_content(ctx, argument)
            elif kind == "EQUATION" and argument and payload:
                _add_equation(ctx.document, int(argument), payload)
            elif kind == "REFERENCES":
                _add_references(ctx)
            elif kind == "PAGE_BREAK":
                ctx.document.add_page_break()
            else:
                raise ValueError(f"Malformed or unknown report token: {line}")
            continue
        paragraph = ctx.document.add_paragraph()
        # Preserve code-like snippets without introducing fragile Markdown dependencies.
        chunks = re.split(r"(`[^`]+`)", line)
        for chunk in chunks:
            if chunk.startswith("`") and chunk.endswith("`"):
                run = paragraph.add_run(chunk[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            else:
                paragraph.add_run(chunk)


def build_report(manifest_path: Path, output_docx: Path) -> Path:
    """Build the detailed DOCX report from one deterministic result manifest."""
    manifest_path = Path(manifest_path).resolve()
    output_docx = Path(output_docx).resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    results_root = manifest_path.parent
    project_root = results_root.parent
    document = Document()
    document.core_properties.title = "Effect of Rear-Wing Angle on F1 Car Performance"
    document.core_properties.subject = "Advanced reproducible motorsport simulation project"
    document.core_properties.author = "ME3106 Project"
    document.core_properties.comments = f"Generated from run {manifest['run_id']}"
    _configure_document(document, manifest["run_id"])
    _add_cover(document, manifest)
    _add_front_matter(document, manifest)
    content = (project_root / "report" / "report_content.md").read_text(encoding="utf-8")
    content = _replace_placeholders(content, manifest)
    context = ReportContext(project_root, results_root, manifest, document)
    _render_markdown(context, content)
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)
    return output_docx


def convert_office_to_pdf(source: Path, output_dir: Path) -> Path:
    """Convert a DOCX or PPTX to PDF using LibreOffice in headless mode."""
    source = Path(source).resolve()
    output_dir = Path(output_dir).resolve()
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        raise RuntimeError("LibreOffice is required for PDF conversion but was not found.")
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="f1wing-libreoffice-") as temporary:
        temporary_root = Path(temporary)
        profile = temporary_root / "profile"
        config = temporary_root / "config"
        cache = temporary_root / "cache"
        runtime = temporary_root / "runtime"
        for directory in (profile, config, cache, runtime):
            directory.mkdir(mode=0o700)
        environment = os.environ.copy()
        environment.update(
            {
                "XDG_CONFIG_HOME": str(config),
                "XDG_CACHE_HOME": str(cache),
                "XDG_RUNTIME_DIR": str(runtime),
            }
        )
        completed = subprocess.run(
            [
                executable,
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(source),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=180,
            env=environment,
        )
    output = output_dir / f"{source.stem}.pdf"
    if completed.returncode != 0 or not output.is_file():
        if source.suffix.lower() == ".docx":
            _render_docx_to_pdf(source, output)
        else:
            raise RuntimeError(
                "LibreOffice conversion failed: "
                + (completed.stderr.strip() or completed.stdout.strip() or "unknown error")
            )
    return output


def _iter_docx_blocks(document: DocumentObject):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield DocxParagraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield DocxTable(child, document)


def _render_docx_to_pdf(source: Path, output: Path) -> None:
    """Portable DOCX fallback renderer used when office conversion is unavailable."""
    document = Document(source)
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "F1Body",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        alignment=TA_JUSTIFY,
        textColor=colors.HexColor("#13233A"),
        spaceAfter=5,
    )
    title = ParagraphStyle(
        "F1Title",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=25,
        leading=30,
        alignment=TA_LEFT,
        textColor=colors.HexColor("#13233A"),
        spaceAfter=14,
    )
    heading1 = ParagraphStyle(
        "F1H1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=21,
        textColor=colors.HexColor("#13233A"),
        spaceAfter=10,
    )
    heading2 = ParagraphStyle(
        "F1H2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#D71920"),
        spaceBefore=7,
        spaceAfter=6,
    )
    caption = ParagraphStyle(
        "F1Caption",
        parent=body,
        fontName="Helvetica-Oblique",
        fontSize=7.5,
        leading=9,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#64748B"),
    )
    equation = ParagraphStyle(
        "F1Equation",
        parent=body,
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        alignment=TA_CENTER,
        spaceBefore=6,
        spaceAfter=6,
    )
    pdf = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        leftMargin=1.8 * cm,
        rightMargin=1.6 * cm,
        topMargin=1.7 * cm,
        bottomMargin=1.6 * cm,
        title="Effect of Rear-Wing Angle on F1 Car Performance",
        author="ME3106 Project",
        pageCompression=1,
    )
    story = []
    buffers: list[io.BytesIO] = []
    first_heading = True
    for block in _iter_docx_blocks(document):
        if isinstance(block, DocxParagraph):
            page_breaks = block._p.xpath('.//w:br[@w:type="page"]')
            if page_breaks:
                story.append(PageBreak())
                continue
            blips = block._p.xpath(".//a:blip")
            if blips:
                for blip in blips:
                    relationship = blip.get(qn("r:embed"))
                    if not relationship:
                        continue
                    part = document.part.related_parts[relationship]
                    buffer = io.BytesIO(part.blob)
                    buffers.append(buffer)
                    image = PdfImage(buffer)
                    scale = min((17.2 * cm) / image.imageWidth, (19.0 * cm) / image.imageHeight)
                    image.drawWidth = image.imageWidth * scale
                    image.drawHeight = image.imageHeight * scale
                    image.hAlign = "CENTER"
                    story.extend((Spacer(1, 4), image, Spacer(1, 5)))
                continue
            text = block.text.strip()
            if not text:
                story.append(Spacer(1, 4))
                continue
            safe = html.escape(text).replace("\n", "<br/>")
            style_name = block.style.name if block.style is not None else "Normal"
            if style_name == "Title":
                chosen = title
            elif style_name == "Subtitle":
                chosen = heading2
            elif style_name == "Heading 1":
                if not first_heading:
                    story.append(PageBreak())
                first_heading = False
                chosen = heading1
            elif style_name == "Heading 2" or style_name == "Heading 3":
                chosen = heading2
            elif style_name == "Caption":
                chosen = caption
            elif style_name == "Equation":
                chosen = equation
            else:
                chosen = body
                if style_name.startswith("List"):
                    safe = "• " + safe
            story.append(PdfParagraph(safe, chosen))
        else:
            rows = []
            for row in block.rows:
                rows.append([PdfParagraph(html.escape(cell.text or " "), body) for cell in row.cells])
            if not rows:
                continue
            count = len(rows[0])
            usable = 17.4 * cm
            if count == 2:
                widths = (0.32 * usable, 0.68 * usable)
            else:
                widths = tuple(usable / count for _ in range(count))
            table = LongTable(rows, colWidths=widths, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#13233A")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
                        ("LEADING", (0, 0), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#94A3B8")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), (colors.white, colors.HexColor("#EAF0F6"))),
                        ("LEFTPADDING", (0, 0), (-1, -1), 3),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                        ("TOPPADDING", (0, 0), (-1, -1), 3),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ]
                )
            )
            story.extend((Spacer(1, 4), table, Spacer(1, 7)))

    def page_decor(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawRightString(A4[0] - 1.6 * cm, A4[1] - 1.0 * cm, "ME3106 · F1 WING-ANGLE PERFORMANCE STUDY")
        canvas.drawCentredString(A4[0] / 2, 0.8 * cm, f"PAGE {doc.page}")
        canvas.restoreState()

    pdf.build(story, onFirstPage=page_decor, onLaterPages=page_decor)
