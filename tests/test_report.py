from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfReader

from f1wing.documents import build_report, convert_office_to_pdf


ROOT = Path(__file__).resolve().parents[1]
REQUIRED_REPORT_HEADINGS = {
    "Abstract",
    "1. Introduction and Research Question",
    "2. Literature Review",
    "3. Aerodynamic Theory",
    "4. Vehicle and Tyre Dynamics",
    "5. Model Architecture and Implementation",
    "6. Verification and Reproducibility",
    "7. Results",
    "8. Uncertainty and Sensitivity",
    "9. Engineering Interpretation",
    "10. Assumptions and Limitations",
    "11. Conclusions and Future Work",
    "References",
    "Appendix A — Parameter Register",
    "Appendix B — MATLAB and Simulink Workflow",
}


def _docx_text(path: Path) -> str:
    document = Document(path)
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


@pytest.fixture(scope="module")
def generated_report(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, Path]:
    output = tmp_path_factory.mktemp("report")
    docx_path = output / "F1_Wing_Angle_Performance_Report.docx"
    build_report(ROOT / "results" / "manifest.json", docx_path)
    pdf_path = convert_office_to_pdf(docx_path, output)
    return docx_path, pdf_path


def test_report_contains_required_sections_and_run_id(generated_report):
    docx_path, _ = generated_report
    manifest = json.loads((ROOT / "results" / "manifest.json").read_text())
    text = _docx_text(docx_path)
    for heading in REQUIRED_REPORT_HEADINGS:
        assert heading in text
    assert manifest["run_id"] in text
    assert "representative educational model" in text.lower()


def test_every_manifest_headline_value_appears_in_report(generated_report):
    docx_path, _ = generated_report
    manifest = json.loads((ROOT / "results" / "manifest.json").read_text())
    text = _docx_text(docx_path)
    for group in manifest["headline_metrics"].values():
        for value in group.values():
            assert f"{value:.3f}" in text


def test_report_includes_all_manifest_figures_and_source_captions(generated_report):
    docx_path, _ = generated_report
    manifest = json.loads((ROOT / "results" / "manifest.json").read_text())
    document = Document(docx_path)
    text = _docx_text(docx_path)
    assert len(document.inline_shapes) >= len(manifest["figures"])
    for index, figure in enumerate(manifest["figures"], start=1):
        assert f"Figure {index}." in text
        assert figure["caption"] in text


def test_report_has_equations_tables_citations_and_page_numbers(generated_report):
    docx_path, _ = generated_report
    document = Document(docx_path)
    text = _docx_text(docx_path)
    assert "Equation (1)" in text and "Equation (12)" in text
    assert len(document.tables) >= 8
    assert "[1]" in text and "[8]" in text
    assert any("PAGE" in run.text for section in document.sections for p in section.footer.paragraphs for run in p.runs)


def test_rendered_report_is_substantive_and_readable(generated_report):
    docx_path, pdf_path = generated_report
    assert docx_path.stat().st_size > 500_000
    assert pdf_path.is_file() and pdf_path.stat().st_size > 500_000
    reader = PdfReader(str(pdf_path))
    assert 25 <= len(reader.pages) <= 60
    sample = "\n".join((reader.pages[i].extract_text() or "") for i in (0, 1, -1))
    assert "Effect of Rear-Wing Angle" in sample
    assert "Appendix" in sample

