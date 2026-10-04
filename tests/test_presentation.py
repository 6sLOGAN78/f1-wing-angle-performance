"""Unit and validation tests for the PowerPoint presentation generator."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation
from pypdf import PdfReader

from f1wing.documents import convert_office_to_pdf
from f1wing.presentation import (
    all_slide_text_within_safe_bounds,
    build_presentation,
    extract_slide_titles,
)

ROOT = Path(__file__).resolve().parents[1]

REQUIRED_PRESENTATION_TITLES = {
    "Simulation and Analysis of the Effect of Rear-Wing Angle on F1 Car Performance",
    "Motivation: Aerodynamic Trade-offs in Modern Formula 1",
    "Research Question and Project Objectives",
    "Aerodynamic Theory: Lift, Induced Drag, and Stall",
    "Vehicle Dynamics: Tyre Load Sensitivity and G-G Ellipse",
    "System Architecture and Simulation Pipeline",
    "Vehicle Parameters and Baseline Configuration",
    "Aerodynamic Characterisation: Polar and Efficiency",
    "Speed-Dependent Aerodynamic Forces and Balance",
    "Straight-Line Performance: Acceleration and Top Speed",
    "Cornering Performance and Lateral Limits",
    "Circuit Archetypes and Fixed-Angle Optimization",
    "Lap-Time Sweeps and Speed Traces",
    "Aerodynamic Strategies: Fixed, Open-Wing, and Active Aero",
    "Uncertainty Analysis and Global Parameter Sensitivity",
    "Model Assumptions, Limitations, and Ethical Disclosure",
    "Conclusions and Engineering Recommendations",
    "References and Deliverable Manifest",
}


@pytest.fixture(scope="module")
def generated_presentation(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, Path]:
    output = tmp_path_factory.mktemp("presentation")
    pptx_path = output / "F1_Wing_Angle_Performance_Presentation.pptx"
    manifest_path = ROOT / "results" / "manifest.json"
    content_path = ROOT / "presentation" / "slide_content.json"
    build_presentation(manifest_path, content_path, pptx_path)
    pdf_path = convert_office_to_pdf(pptx_path, output)
    return pptx_path, pdf_path


def _deck_text(deck: Presentation) -> str:
    parts = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        parts.append(cell.text)
    return "\n".join(parts)


def test_deck_has_expected_structure_and_no_overflow(generated_presentation):
    pptx_path, _ = generated_presentation
    deck = Presentation(str(pptx_path))
    assert 16 <= len(deck.slides) <= 20
    assert all_slide_text_within_safe_bounds(deck)
    titles = extract_slide_titles(deck)
    for req in REQUIRED_PRESENTATION_TITLES:
        assert any(req.lower() in t.lower() for t in titles), f"Missing required slide title matching '{req}'"


def test_every_manifest_headline_value_appears_in_presentation(generated_presentation):
    pptx_path, _ = generated_presentation
    manifest = json.loads((ROOT / "results" / "manifest.json").read_text(encoding="utf-8"))
    deck = Presentation(str(pptx_path))
    text = _deck_text(deck)
    for group_name, group in manifest["headline_metrics"].items():
        for circuit_name, value in group.items():
            formatted = f"{value:.3f}"
            assert formatted in text, f"Headline metric {group_name}.{circuit_name} ({formatted}) missing from deck"


def test_presentation_contains_run_id_and_classification(generated_presentation):
    pptx_path, _ = generated_presentation
    manifest = json.loads((ROOT / "results" / "manifest.json").read_text(encoding="utf-8"))
    deck = Presentation(str(pptx_path))
    text = _deck_text(deck)
    assert manifest["run_id"] in text
    assert "representative educational model" in text.lower()


def test_presentation_has_embedded_figures_and_tables(generated_presentation):
    pptx_path, _ = generated_presentation
    deck = Presentation(str(pptx_path))
    picture_count = sum(
        1 for slide in deck.slides for shape in slide.shapes if shape.shape_type == 13  # MSO_SHAPE_TYPE.PICTURE
    )
    table_count = sum(1 for slide in deck.slides for shape in slide.shapes if shape.has_table)
    assert picture_count >= 8
    assert table_count >= 4


def test_rendered_presentation_pdf_is_valid(generated_presentation):
    pptx_path, pdf_path = generated_presentation
    assert pptx_path.is_file() and pptx_path.stat().st_size > 50_000
    assert pdf_path.is_file() and pdf_path.stat().st_size > 50_000
    reader = PdfReader(str(pdf_path))
    assert 16 <= len(reader.pages) <= 20
